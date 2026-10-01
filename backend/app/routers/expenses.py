import datetime
import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Form, File, UploadFile, Request
from sqlalchemy.orm import Session
from typing import List
from .. import models, schemas, file_storage, reporting
from ..database import get_db
from ..security import current_user

router = APIRouter(prefix="/expenses", tags=["Expenses"])

ALLOWED_IMAGE_TYPES = ("image/jpeg", "image/png", "image/webp", "image/heic", "image/heif")
MAX_PHOTO_BYTES = 5 * 1024 * 1024


def _check_expense_date(new_date: datetime.datetime, user, current: datetime.datetime = None):
    """
    Back-dating rule: only administrators may put an expense on a day other than
    today (Pakistan time). Nobody may date one in the future. When editing, only a
    genuine change of day is checked, so ordinary edits by non-admins keep working.
    """
    today = reporting.pkt_today()
    day = reporting.pkt_date(new_date)
    if current is not None and reporting.pkt_date(current) == day:
        return
    if day > today:
        raise HTTPException(status_code=400, detail="An expense can't be dated in the future")
    if day != today and not user.role.is_admin:
        raise HTTPException(status_code=403, detail="Only an administrator can enter or change an expense to a past date")


@router.post("/", response_model=schemas.ExpenseOut)
async def create_expense(
    request: Request,
    data: str = Form(..., description="JSON-encoded ExpenseCreate payload"),
    bill_photo: UploadFile = File(..., description="Photo of the receipt/bill — required"),
    db: Session = Depends(get_db),
):
    """A bill photo is compulsory for every expense, saved to the server's
    persistent volume (see app/file_storage.py)."""
    try:
        expense = schemas.ExpenseCreate.model_validate_json(data)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Invalid data payload: {e}")

    if bill_photo.content_type not in ALLOWED_IMAGE_TYPES:
        raise HTTPException(status_code=400, detail="Bill photo must be a JPEG, PNG, WEBP, or HEIC image")
    photo_bytes = await bill_photo.read()
    if not photo_bytes:
        raise HTTPException(status_code=400, detail="Bill photo file is empty")
    if len(photo_bytes) > MAX_PHOTO_BYTES:
        raise HTTPException(status_code=400, detail="Bill photo must be under 5MB")

    payload = expense.dict()
    if not payload.get("expense_date"):
        payload["expense_date"] = datetime.datetime.utcnow()
    else:
        _check_expense_date(payload["expense_date"], current_user(request))
        if reporting.pkt_date(payload["expense_date"]) == reporting.pkt_today():
            payload["expense_date"] = datetime.datetime.utcnow()   # today: keep the real time of entry

    try:
        saved = file_storage.save_bill_photo(photo_bytes, bill_photo.filename)
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Couldn't save the bill photo: {e}")

    payload["bill_photo_url"] = saved["url"]

    db_expense = models.Expense(**payload)
    db.add(db_expense)
    db.commit()
    db.refresh(db_expense)
    return db_expense


@router.get("/report")
def expense_report(request: Request, start_date: datetime.date = None, end_date: datetime.date = None,
                   db: Session = Depends(get_db)):
    """
    Excel report of every expense in the date range (inclusive, Pakistan time):
    one line per expense with its full detail, a total after each day, and a gross
    total at the end.
    """
    lo, hi = reporting.range_bounds(start_date, end_date)
    rows = db.query(models.Expense).filter(
        models.Expense.expense_date >= lo, models.Expense.expense_date < hi,
    ).order_by(models.Expense.expense_date, models.Expense.id).all()

    by_day = {}
    for e in rows:
        by_day.setdefault(reporting.pkt_date(e.expense_date), []).append(e)

    user = current_user(request)
    wb = openpyxl.Workbook()
    columns = ["Date", "Category", "Cost Center", "Description", "Notes", "Paid Via", "Payment Details", "Amount (Rs.)"]
    ws, r = reporting.new_sheet(
        wb, "Expenses", "Riwayat Oils and Fats — Expense Report",
        [f"Period: {reporting.fmt_day(start_date)} to {reporting.fmt_day(end_date)}",
         f"Generated {reporting.fmt_day(reporting.pkt_today())} by {user.full_name}"],
        columns, [14, 14, 16, 34, 30, 11, 38, 16],
    )
    n = len(columns)
    gross = 0.0
    if not rows:
        ws.cell(row=r, column=1, value="No expenses were recorded in this period.")
        r += 1
    for day in sorted(by_day):
        day_total = 0.0
        for e in by_day[day]:
            method = getattr(e.payment_method, "value", e.payment_method)
            reporting.write_row(ws, r, [
                day, e.category.value.title(), e.plant.value.title() if e.plant else "General / Admin",
                e.description or "", e.notes or "", (method or "—").title(),
                reporting.describe_payment(e) if method else "—", e.amount,
            ], money_cols=(8,), wrap_cols=(4, 5, 7))
            ws.cell(row=r, column=1).number_format = "dd mmm yyyy"
            day_total += e.amount
            r += 1
        reporting.write_total_row(ws, r, f"Total for {reporting.fmt_day(day)}  ({len(by_day[day])} expense{'s' if len(by_day[day]) != 1 else ''})",
                                  1, {8: day_total}, n, reporting.DAY_FILL)
        gross += day_total
        r += 2
    reporting.write_total_row(ws, r, f"GROSS TOTAL  ({len(rows)} expense{'s' if len(rows) != 1 else ''}, {len(by_day)} day{'s' if len(by_day) != 1 else ''})",
                              1, {8: gross}, n, reporting.GROSS_FILL, border=reporting.TOP_DOUBLE)
    return reporting.workbook_response(wb, f"expenses_{start_date}_to_{end_date}.xlsx")


@router.get("/", response_model=List[schemas.ExpenseOut])
def list_expenses(plant: str = None, category: str = None, db: Session = Depends(get_db)):
    query = db.query(models.Expense)
    if plant:
        query = query.filter(models.Expense.plant == plant)
    if category:
        query = query.filter(models.Expense.category == category)
    return query.order_by(models.Expense.expense_date.desc()).all()


@router.put("/{expense_id}", response_model=schemas.ExpenseOut)
def update_expense(expense_id: int, update: schemas.ExpenseUpdate, request: Request, db: Session = Depends(get_db)):
    expense = db.query(models.Expense).filter(models.Expense.id == expense_id).first()
    if not expense:
        raise HTTPException(status_code=404, detail="Expense not found")
    if update.expense_date is not None:
        _check_expense_date(update.expense_date, current_user(request), current=expense.expense_date)
    for field, value in update.dict(exclude_unset=True).items():
        setattr(expense, field, value)
    db.commit()
    db.refresh(expense)
    return expense


@router.delete("/{expense_id}")
def delete_expense(expense_id: int, db: Session = Depends(get_db)):
    expense = db.query(models.Expense).filter(models.Expense.id == expense_id).first()
    if not expense:
        raise HTTPException(status_code=404, detail="Expense not found")
    db.delete(expense)
    db.commit()
    return {"message": "Expense deleted", "expense_id": expense_id}
