"""Shared helpers for the Excel report downloads (expenses, purchase orders)."""
import io
import datetime
from typing import Optional

import openpyxl
import openpyxl.worksheet.properties
from fastapi import HTTPException
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# The business runs on Pakistan time; timestamps are stored as naive UTC.
PKT_OFFSET = datetime.timedelta(hours=5)
MONEY_FMT = '#,##0.00'
QTY_FMT = '#,##0.###'

TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
DAY_FILL = PatternFill("solid", fgColor="DDE6F0")
GROSS_FILL = PatternFill("solid", fgColor="F3D9A4")
THIN = Side(style="thin", color="B7C0CC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP_DOUBLE = Border(top=Side(style="double"), bottom=Side(style="double"), left=THIN, right=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def pkt_date(dt: datetime.datetime) -> datetime.date:
    """The Pakistan calendar day a stored (UTC) timestamp falls on."""
    return (dt + PKT_OFFSET).date()


def pkt_today() -> datetime.date:
    return pkt_date(datetime.datetime.utcnow())


def range_bounds(start: datetime.date, end: datetime.date):
    """[start 00:00 PKT, end+1 00:00 PKT) expressed in naive UTC for DB filtering."""
    if start is None or end is None:
        raise HTTPException(status_code=400, detail="Choose both a From date and a To date")
    if start > end:
        raise HTTPException(status_code=400, detail="The From date must be on or before the To date")
    lo = datetime.datetime.combine(start, datetime.time.min) - PKT_OFFSET
    hi = datetime.datetime.combine(end + datetime.timedelta(days=1), datetime.time.min) - PKT_OFFSET
    return lo, hi


def fmt_day(d: datetime.date) -> str:
    return d.strftime("%d %b %Y")


def describe_payment(entry, vendor_account=None) -> str:
    """One-line human description of how something was paid."""
    method = getattr(entry.payment_method, "value", entry.payment_method)
    if not method:
        return "—"
    if method == "cash":
        return "Cash"
    if method == "cheque":
        return f"Cheque #{entry.cheque_number or '—'} drawn on {entry.cheque_bank or '—'}"
    if method == "online":
        src = f" from {entry.our_bank}" if entry.our_bank else ""
        if vendor_account is not None:
            title = f" – {vendor_account.account_title}" if vendor_account.account_title else ""
            dest = f"{vendor_account.bank_name}{title} – A/C {vendor_account.account_number}"
        else:
            dest = f"{entry.other_party_name or '—'} ({entry.other_party_bank or '—'})"
        return f"Online transfer{src} to {dest}"
    return str(method)


def new_sheet(wb, title: str, heading: str, sub_lines, columns, widths):
    """Adds a sheet with a title block and a styled header row. Returns (ws, next_row)."""
    ws = wb.create_sheet(title)
    ws["A1"] = heading
    ws["A1"].font = TITLE_FONT
    r = 2
    for line in sub_lines:
        ws.cell(row=r, column=1, value=line).font = Font(italic=True, color="555555")
        r += 1
    r += 1
    for c, name in enumerate(columns, start=1):
        cell = ws.cell(row=r, column=c, value=name)
        cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[r].height = 30
    ws.freeze_panes = ws.cell(row=r + 1, column=1)
    # print-friendly: landscape, one page wide, header row repeated on every page
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    ws.print_title_rows = f"{r}:{r}"
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws, r + 1


def write_row(ws, r, values, money_cols=(), qty_cols=(), wrap_cols=()):
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.border = BORDER
        cell.alignment = WRAP_TOP if c in wrap_cols else Alignment(vertical="top")
        if c in money_cols:
            cell.number_format = MONEY_FMT
        if c in qty_cols:
            cell.number_format = QTY_FMT


def write_total_row(ws, r, label, label_col, total_cols: dict, ncols, fill, border=BORDER):
    """A bold, shaded total line. total_cols = {column_index: number}."""
    for c in range(1, ncols + 1):
        cell = ws.cell(row=r, column=c)
        cell.fill, cell.font, cell.border = fill, BOLD, border
    ws.cell(row=r, column=label_col, value=label)
    for c, v in total_cols.items():
        cell = ws.cell(row=r, column=c, value=v)
        cell.number_format = MONEY_FMT


def workbook_response(wb, filename: str) -> StreamingResponse:
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        del wb["Sheet"]
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
